# -*- coding: utf-8 -*-
import sys; sys.path.insert(0,'.')
from datos_marketing import CAT
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.utils import get_column_letter as L
OUT=sys.argv[1]
ORIG={'Pauta en redes sociales':'Marketing S/700 al mes','Degustaciones extra':'Marketing S/700 al mes','Material de stand y POP':'Marketing S/700 al mes + lanzamiento',
 'Mamás embajadoras (premio al padre)':'Embajadoras S/200 al mes','Pauta de expectativa':'Lanzamiento (inversión)','Video de lanzamiento con niños':'Lanzamiento (inversión)',
 'Muestras de degustación (600)':'Lanzamiento (inversión)','Influencers (40 packs)':'Lanzamiento (inversión)','Vendedor(a) en ferias pagadas':'Ferias pagadas',
 'Ferias pagadas (stand)':'Ferias pagadas','Vendedor(a) en eventos sin costo':'Eventos sin costo','Aporte a eventos sin costo':'Eventos sin costo'}
F='Arial'; W='FFFFFF'
HF=PatternFill('solid',fgColor='4A2018'); CF=PatternFill('solid',fgColor='E8CDB3'); TF=PatternFill('solid',fgColor='BA4036'); ZF=PatternFill('solid',fgColor='FBF3E8')
thin=Side(style='thin',color='BFBFBF'); BX=Border(top=thin,bottom=thin,left=thin,right=thin)
SOL='#,##0;(#,##0);-'
wb=Workbook(); ws=wb.active; ws.title='Presupuesto de marketing'
def c(ref,v,bold=False,color='000000',fill=None,fmt=None,al=None,size=10):
    x=ws[ref]; x.value=v; x.font=Font(name=F,bold=bold,color=color,size=size)
    if fill: x.fill=fill
    if fmt: x.number_format=fmt
    x.alignment=Alignment(horizontal=al or ('right' if fmt else 'left'),vertical='center',wrap_text=True); x.border=BX
ws['B2'].value='AndiBite — Presupuesto de marketing por medio (S/ sin IGV), prelanzamiento y año 1'; ws['B2'].font=Font(name=F,bold=True,size=14,color='4A2018')
ws['B3'].value='Mismas cifras del Excel v7 y del documento 10. Celdas azules = datos; el resto son fórmulas. Los stands (espacio y vendedores de fin de semana) van en el costo de venta, no aquí.'; ws['B3'].font=Font(name=F,size=9,italic=True)
MES=['Prelanz. oct-dic 2026','Ene','Feb','Mar','Abr','May','Jun','Jul','Ago','Set','Oct','Nov','Dic']
hdr=['Medio','Dónde o con quién','Meses activos','Presupuesto mensual (promedio)']+MES+['Presupuesto total','Sale de la línea del Excel v7']
for j,h in enumerate(hdr): c(f'{L(2+j)}5',h,True,W,HF,al='center')
ws.row_dimensions[5].height=42
r=6; catrows=[]; itemrows=[]
c0,c1=L(6),L(18)  # F..R meses
for cat,items in CAT:
    cr=r; catrows.append(cr); r+=1; first=r
    for name,lugar,vals in items:
        c(f'B{r}',name); c(f'C{r}',lugar)
        for j,v in enumerate(vals):
            x=ws[f'{L(6+j)}{r}']; x.value=v; x.font=Font(name=F,color='0000FF',size=10); x.number_format=SOL; x.border=BX
        c(f'D{r}',f'=COUNTIF(G{r}:R{r},">0")+IF(F{r}>0,1,0)',fmt='0',al='center')
        c(f'E{r}',f'=IF(D{r}=0,0,S{r}/D{r})',fmt=SOL)
        c(f'S{r}',f'=SUM(F{r}:R{r})',True,fmt=SOL)
        c(f'T{r}',ORIG[name],size=9)
        itemrows.append(r); r+=1
    last=r-1
    c(f'B{cr}',cat,True,'4A2018',CF)
    for col in ['C','D','E','T']: c(f'{col}{cr}',None,fill=CF)
    for j in range(13): col=L(6+j); c(f'{col}{cr}',f'=SUM({col}{first}:{col}{last})',True,fill=CF,fmt=SOL)
    c(f'S{cr}',f'=SUM(S{first}:S{last})',True,fill=CF,fmt=SOL)
tr=r; c(f'B{tr}','PRESUPUESTO TOTAL',True,W,TF)
for col in ['C','D','E','T']: c(f'{col}{tr}',None,fill=TF)
for j in range(14):
    col=L(6+j); c(f'{col}{tr}','='+'+'.join(f'{col}{x}' for x in catrows),True,W,TF,fmt=SOL)
# conciliación con el Excel v7
q=tr+2
ws[f'B{q}'].value='De dónde sale cada sol (conciliación con el Excel v7)'; ws[f'B{q}'].font=Font(name=F,bold=True,size=11,color='4A2018')
rows=[('Marketing fijo: pauta, degustación y material (S/700 al mes)','Costos fijos, hoja Supuestos',8400),
      ('Mamás embajadoras (S/200 al mes)','Costos fijos, hoja Supuestos',2400),
      ('Total: presupuesto de marketing de S/900 al mes','12 meses de 2027',10800),
      ('Marketing de lanzamiento','Inversión inicial, hoja Inversion',3204),
      ('Ferias pagadas: stand y vendedor(a) 3 días (febrero y abril)','Hoja Mensual, ferias pagadas',4140),
      ('Eventos sin costo: vendedor(a) S/90 y aporte S/70 (40 eventos)','Hoja Mensual, eventos sin costo',6400),
      ('Total del plan de marketing','Debe ser igual al presupuesto total de arriba',24544)]
for i,(a,b,v) in enumerate(rows):
    rr=q+1+i; bold=a.startswith('Total')
    c(f'B{rr}',a,bold); ws.merge_cells(f'B{rr}:E{rr}'); c(f'F{rr}',b,size=9); ws.merge_cells(f'F{rr}:K{rr}')
    x=ws[f'L{rr}']; x.value=v; x.font=Font(name=F,bold=bold,color='0000FF' if not bold else '000000'); x.number_format=SOL; x.border=BX
ws[f'L{q+3}'].value=f'=L{q+1}+L{q+2}'; ws[f'L{q+7}'].value=f'=L{q+3}+L{q+4}+L{q+5}+L{q+6}'
c(f'M{q+7}',f'=IF(L{q+7}=S{tr},"Cuadra","Revisar")',True,al='center')
# resumen por tipo de medio (para el gráfico)
s0=q+10
ws[f'B{s0}'].value='Presupuesto por tipo de medio'; ws[f'B{s0}'].font=Font(name=F,bold=True,size=11,color='4A2018')
for i,cr in enumerate(catrows):
    ws[f'B{s0+1+i}'].value=f'=B{cr}'; ws[f'C{s0+1+i}'].value=f'=S{cr}'; ws[f'C{s0+1+i}'].number_format=SOL
    ws[f'D{s0+1+i}'].value=f'=C{s0+1+i}/S{tr}'; ws[f'D{s0+1+i}'].number_format='0%'
bar=BarChart(); bar.type='col'; bar.title='Presupuesto por tipo de medio'; bar.style=10; bar.legend=None
bar.add_data(Reference(ws,min_col=3,min_row=s0+1,max_row=s0+4)); bar.set_categories(Reference(ws,min_col=2,min_row=s0+1,max_row=s0+4))
bar.dataLabels=DataLabelList(); bar.dataLabels.showVal=True; bar.dataLabels.showSerName=False; bar.dataLabels.showCatName=False; bar.dataLabels.showLegendKey=False; bar.dataLabels.showPercent=False; bar.series[0].graphicalProperties.solidFill='4A2018'
bar.height=8; bar.width=14; ws.add_chart(bar,f'B{s0+6}')
line=LineChart(); line.title='Presupuesto de marketing mensual'; line.style=12; line.legend=None
line.add_data(Reference(ws,min_col=6,max_col=18,min_row=tr),from_rows=True); line.set_categories(Reference(ws,min_col=6,max_col=18,min_row=5))
line.series[0].graphicalProperties.line.solidFill='BA4036'; line.series[0].graphicalProperties.line.width=28000; line.series[0].smooth=False; line.series[0].marker.symbol='circle'; line.series[0].marker.size=7
line.height=8; line.width=24; ws.add_chart(line,f'H{s0+6}')
widths={'A':2,'B':34,'C':28,'D':9,'E':12,'S':12,'T':30}
for k,v in widths.items(): ws.column_dimensions[k].width=v
for j in range(13): ws.column_dimensions[L(6+j)].width=9 if j else 13
ws.freeze_panes='F6'
ws.page_setup.orientation='landscape'; ws.sheet_properties.pageSetUpPr.fitToPage=True
wb.save(OUT); print('ok',OUT,'total row',tr)
