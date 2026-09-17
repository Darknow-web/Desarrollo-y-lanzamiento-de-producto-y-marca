# -*- coding: utf-8 -*-
"""Flat, high-fidelity rebuild of the AndiBite pouch design at exact bag sizes.
Layout is specified once (in cm) and rendered twice: PIL (print PNG/PDF) and python-pptx (editable in Canva)."""
from PIL import Image, ImageDraw, ImageFont, ImageFilter, ImageChops
import math, os
DPI=300
X4=Image.open('../sr/mockup_x4.png').convert('RGB')
KRAFT=(232,205,179); BAND=(90,36,33); RIBBON=(112,42,40); BADGE=(99,39,39); BROWN=(92,48,30); RED=(186,64,54); CREMA=(247,236,218)
OLIVE=(135,125,52)
FONTS={'fredoka':'fonts/FredokaOne.ttf','caveat':'fonts/Caveat-Bold.ttf','nunito':'fonts/Nunito-Regular.ttf','nunitob':'fonts/Nunito-Bold.ttf','nunitox':'fonts/Nunito-ExtraBold.ttf'}
PPTX_FONT={'fredoka':'Fredoka One','caveat':'Caveat','nunito':'Nunito','nunitob':'Nunito','nunitox':'Nunito'}
def px(c): return int(round(c/2.54*DPI))
_FC={}
def _f(fname,size_cm):
    k=(fname,px(size_cm))
    if k not in _FC: _FC[k]=ImageFont.truetype(FONTS[fname],max(6,px(size_cm)))
    return _FC[k]
def fit(fname,txt,target_w,max_size):
    """largest size (cm) whose widest line fits target_w cm, capped at max_size"""
    lines=txt.split('\n'); size=max_size
    for _ in range(12):
        f=_f(fname,size); w=max(f.getlength(l) for l in lines)/DPI*2.54
        if w<=target_w*1.005: break
        size=size*target_w/w
    return size

# ---------- asset extraction from the upscaled mockup (orig-pixel boxes, x4 inside) ----------
def crop4(box): x0,y0,x1,y1=box; return X4.crop((x0*4,y0*4,x1*4,y1*4))
def floodfill_corners(img, ref_bottom=None):
    w,h=img.size; px_=img.load(); ref=ref_bottom or px_[w//2,h-10]
    for (x,y) in [(0,h-1),(w-1,h-1),(0,h-60),(w-1,h-60)]: ImageDraw.floodfill(img,(x,y),ref,thresh=45)
    return img
def with_fade_top(img, rows):
    """RGBA with alpha fading in over the first `rows` px (blend into flat kraft)"""
    a=Image.new('L',img.size,255); d=ImageDraw.Draw(a)
    for i in range(rows): d.line((0,i,img.width,i),fill=int(255*i/rows))
    img=img.convert('RGBA'); img.putalpha(a); return img
def disc(img, feather=6):
    a=Image.new('L',img.size,0); d=ImageDraw.Draw(a); d.ellipse((0,0,img.width-1,img.height-1),fill=255)
    a=a.filter(ImageFilter.GaussianBlur(feather)); img=img.convert('RGBA'); img.putalpha(a); return img
def keyed(img, bg, thresh=30, feather=3):
    """alpha = pixels that differ from the flat background colour (for the mountain logo / hearts)"""
    diff=ImageChops.difference(img.convert('RGB'), Image.new('RGB',img.size,bg)).convert('L')
    a=diff.point(lambda v:255 if v>thresh else 0).filter(ImageFilter.MaxFilter(5)).filter(ImageFilter.GaussianBlur(feather))
    img=img.convert('RGBA'); img.putalpha(a); return img

def make_assets():
    A='v3/assets/'
    illu=crop4((152,452,624,712)); illu=with_fade_top(illu,60); illu.save(A+'illu_front.png')
    land=floodfill_corners(crop4((795,594,1245,836))); land=with_fade_top(land,40); land.save(A+'landscape_back.png')
    keyed(crop4((312,190,466,262)),KRAFT).save(A+'logo_mountain.png')
    keyed(crop4((216,232,246,260)),KRAFT).save(A+'heart_front.png')
    keyed(crop4((990,286,1032,318)),KRAFT).save(A+'heart_back.png')
    for i,(cx,cy) in enumerate([(231,768),(330,768),(438,768)]): disc(crop4((cx-23,cy-23,cx+23,cy+23)),feather=4).save(A+f'icon_front_{i}.png')
    for i,(cx,cy) in enumerate([(878,492),(1013,492),(1148,492)]): disc(crop4((cx-27,cy-27,cx+27,cy+27)),feather=2).save(A+f'icon_back_{i}.png')
    # front badge glyph area is rebuilt with text; landscape strip for bag 2 comes from landscape_back

# ---------- layout spec (all units cm) ----------
# front reference bag: orig x 150..626 (476), y 118..838 (720). u=(x-150)/476, v=(y-118)/720
def spec_front(W,H,window=None):
    """window=(x,y,w,h) cm or None. Returns list of elements."""
    sx=W/476.0; sy=H/720.0 if window is None else sx   # bag 2 keeps natural proportions
    X=lambda x:(x-150)*sx; Y=lambda y:(y-118)*sy; S=lambda n:n*sx
    E=[('rect',0,0,W,H,KRAFT)]
    # handwritten block (rotated, rising to the right)
    hw="Lo bueno\ntambién\npuede ser\ndelicioso"; sz=fit('caveat',hw,S(84),S(40))
    E.append(('text',hw,X(178),Y(128),S(120),S(110),'caveat',sz,BROWN,'left',8,S(25)))
    E.append(('line',X(157),Y(213),X(169),Y(206),BROWN,S(2.2))); E.append(('line',X(161),Y(200),X(171),Y(195),BROWN,S(2.2)))
    E.append(('line',X(277),Y(173),X(289),Y(166),BROWN,S(2.2))); E.append(('line',X(280),Y(186),X(292),Y(183),BROWN,S(2.2)))
    E.append(('image','v3/assets/heart_front.png',X(221),Y(238),S(26),S(24)))
    # badge tab top-right
    E.append(('roundrect',X(541),-S(20),S(85),Y(196)+S(20),BADGE,S(16)))
    E.append(('text',"6",X(541),Y(121),S(85),S(40),'fredoka',fit('fredoka',"6",S(21),S(40)),CREMA,'center',0,None))
    E.append(('text',"mini\nbrownies",X(541),Y(157),S(85),S(40),'nunitob',fit('nunitob',"mini\nbrownies",S(66),S(20)),CREMA,'center',0,S(16)))
    # logo + name
    E.append(('image','v3/assets/logo_mountain.png',X(312),Y(190),S(154),S(72)))
    E.append(('text2',"Andi","Bite",X(150),Y(266),S(476),S(100),'fredoka',fit('fredoka',"AndiBite",S(410),S(140)),BROWN,RED))
    E.append(('text',"Mini Brownie Nutritivo",X(150),Y(363),S(476),S(36),'nunitox',fit('nunitox',"Mini Brownie Nutritivo",S(328),S(40)),BROWN,'center',0,None))
    # ribbon
    E.append(('roundrect',X(238),Y(403),S(326),S(47),RIBBON,S(14)))
    E.append(('poly',[(X(228),Y(409)),(X(248),Y(403)),(X(248),Y(450)),(X(228),Y(446)),(X(236),Y(427))],(84,28,30)))
    E.append(('poly',[(X(574),Y(409)),(X(554),Y(403)),(X(554),Y(450)),(X(574),Y(446)),(X(566),Y(427))],(84,28,30)))
    E.append(('text',"Sangrecita + Cañihua + Cacao",X(238),Y(412),S(326),S(30),'nunitob',fit('nunitob',"Sangrecita + Cañihua + Cacao",S(312),S(30)),CREMA,'center',0,None))
    if window is None:
        E.append(('image','v3/assets/illu_front.png',0,Y(452),W,Y(712)-Y(452)))
        band_top=Y(706)
    else:
        wx,wy,ww,wh=window
        band_h=(838-706)*sx; band_top=H-band_h
        land_top=Y(452)+S(10)
        E.append(('image','v3/assets/landscape_back.png',0,land_top,W,band_top-land_top+0.05))
        E.append(('window',wx,wy,ww,wh))
    E.append(('rect',0,band_top,W,H-band_top,BAND))
    V=lambda n:n*sy
    cy=band_top+V(62)
    for i,cx in enumerate([231,330,438]):
        E.append(('image',f'v3/assets/icon_front_{i}.png',X(cx)-S(23),cy-S(23),S(46),S(46)))
    labels=[("Fuente\nde hierro",62),("Buena fuente\nde proteína",90),("Con ingredientes\nandinos",108)]
    for cx,(l,tw) in zip([231,330,438],labels):
        E.append(('text',l,X(cx)-S(60),band_top+V(92),S(120),S(50),'nunitob',fit('nunitob',l,S(tw),V(17)),CREMA,'center',0,V(21)))
    E.append(('text',"Peso neto",X(530),band_top+V(66),S(90),S(20),'nunito',fit('nunito',"Peso neto",S(46),V(15)),CREMA,'center',0,None))
    E.append(('text',"120 g",X(530),band_top+V(86),S(90),S(30),'nunitox',fit('nunitox',"120 g",S(50),V(23)),CREMA,'center',0,None))
    return E

# back reference bag: orig x 792..1248 (456), y 150..838 (688)
def spec_back(W,H,legal=True,natural=False):
    sx=W/456.0; sy=H/688.0 if not natural else sx
    X=lambda x:(x-792)*sx; Y=lambda y:(y-150)*sy; S=lambda n:n*sx
    E=[('rect',0,0,W,H,KRAFT)]
    E.append(('text',"De nuestra tierra,",X(792),Y(205),S(456),S(48),'caveat',fit('caveat',"De nuestra tierra,",S(230),S(60)),BROWN,'center',3,None))
    E.append(('text',"a su lonchera",X(792),Y(250),S(456),S(48),'caveat',fit('caveat',"a su lonchera",S(165),S(60)),BROWN,'center',3,None))
    for (x0,y0,x1,y1) in [(884,238,893,231),(881,252,892,250),(885,272,894,264),(1128,232,1137,224),(1132,248,1142,246),(1131,268,1140,260)]:
        E.append(('line',X(x0),Y(y0),X(x1),Y(y1),BROWN,S(2.2)))
    E.append(('image','v3/assets/heart_back.png',X(990),Y(287),S(42),S(32)))
    para="Un mini brownie nutritivo, elaborado\ncon ingredientes peruanos, para que\ncada bocado sea una fuente de energía\ny sabor."
    V=lambda n:n*sy
    psz=fit('nunito',para,S(342),V(27))
    E.append(('text',para,X(792),Y(332),S(456),S(160),'nunito',psz,BROWN,'center',0,psz*1.42))
    for i,cx in enumerate([878,1013,1148]):
        E.append(('image',f'v3/assets/icon_back_{i}.png',X(cx)-S(27),Y(492)-S(27),S(54),S(54)))
    for x in (945,1082): E.append(('line',X(x),Y(462),X(x),Y(562),(205,182,152),S(1.5)))
    labels=[("Ingredientes\nperuanos",114),("Nutrición\nen cada bocado",132),("Hecho\ncon amor",112)]
    for cx,(l,tw) in zip([878,1013,1148],labels):
        E.append(('text',l,X(cx)-S(70),Y(526),S(140),S(60),'nunitob',fit('nunitob',l,S(tw),V(20)),BROWN,'center',0,V(25)))
    if legal:
        lg="Ingredientes: cacao, avena, cañihua, sangrecita liofilizada, plátano, huevo, aceite vegetal, panela. Contiene huevo.\n6 unidades x 20 g  ·  Elaborado: ____/____/2026  ·  Consumir dentro de 5 días  ·  Hecho en Lima  ·  Prototipo de evaluación, no destinado a la venta."
        lsz=fit('nunito',lg,S(430),V(9))
        E.append(('text',lg,X(792)+S(8),Y(160),S(440),S(30),'nunito',lsz,BROWN,'center',0,lsz*1.3))
    E.append(('image','v3/assets/landscape_back.png',0,Y(594),W,H-Y(594)))
    return E

# ---------- PIL renderer ----------
def font(name,size_cm): return _f(name,size_cm)
def draw_text(img,txt,x,y,w,h,fname,size,fill,align,rot,lh):
    f=font(fname,size); lines=txt.split('\n'); lhp=px(lh) if lh else int(f.size*1.15)
    top=min(f.getbbox(l)[1] for l in lines)          # distance from draw origin to the glyph top (cap-top alignment)
    tw=max(f.getlength(l) for l in lines); th=lhp*(len(lines)-1)+f.size
    pad=int(f.size*0.6); layer=Image.new('RGBA',(int(tw)+2*pad,th+2*pad),(0,0,0,0)); d=ImageDraw.Draw(layer)
    for i,l in enumerate(lines):
        lw=f.getlength(l); ox=pad if align=='left' else pad+(tw-lw)/2
        d.text((ox,pad+i*lhp-top),l,font=f,fill=fill)
    if rot:
        layer=layer.rotate(rot,resample=Image.BICUBIC,expand=True)
    if align=='left': bx=px(x)-pad
    else: bx=px(x)+(px(w)-layer.width)//2
    img.paste(layer,(int(bx),px(y)-pad),layer)
def render_pil(E,W,H,out):
    img=Image.new('RGB',(px(W),px(H)),KRAFT); d=ImageDraw.Draw(img)
    for e in E:
        k=e[0]
        if k=='rect': _,x,y,w,h,c=e; d.rectangle((px(x),px(y),px(x+w),px(y+h)),fill=c)
        elif k=='roundrect': _,x,y,w,h,c,r=e; d.rounded_rectangle((px(x),px(y),px(x+w),px(y+h)),radius=px(r),fill=c)
        elif k=='poly': d.polygon([(px(a),px(b)) for a,b in e[1]],fill=e[2])
        elif k=='line': _,x0,y0,x1,y1,c,wd=e; d.line((px(x0),px(y0),px(x1),px(y1)),fill=c,width=max(2,px(wd)))
        elif k=='image':
            _,p,x,y,w,h=e; im=Image.open(p).convert('RGBA').resize((max(1,px(w)),max(1,px(h))),Image.LANCZOS); img.paste(im,(px(x),px(y)),im)
        elif k=='text': draw_text(img,*e[1:])
        elif k=='text2':
            _,a,b,x,y,w,h,fname,size,ca,cb=e; f=font(fname,size); wa=f.getlength(a); wb=f.getlength(b); x0=px(x)+(px(w)-wa-wb)/2; top=f.getbbox(a+b)[1]
            d.text((x0,px(y)-top),a,font=f,fill=ca); d.text((x0+wa,px(y)-top),b,font=f,fill=cb)
        elif k=='window':
            _,x,y,w,h=e; d.rectangle((px(x),px(y),px(x+w),px(y+h)),fill='white')
            for xx in range(px(x),px(x+w),24): d.line((xx,px(y),min(xx+12,px(x+w)),px(y)),fill=(120,120,120),width=2); d.line((xx,px(y+h),min(xx+12,px(x+w)),px(y+h)),fill=(120,120,120),width=2)
            for yy in range(px(y),px(y+h),24): d.line((px(x),yy,px(x),min(yy+12,px(y+h))),fill=(120,120,120),width=2); d.line((px(x+w),yy,px(x+w),min(yy+12,px(y+h))),fill=(120,120,120),width=2)
            f=font('nunito',0.25); t='VENTANA: recortar antes de pegar'; tw=f.getlength(t); d.text((px(x)+(px(w)-tw)/2,px(y)+px(h)/2-f.size/2),t,font=f,fill=(160,160,160))
    img.save(out,dpi=(DPI,DPI)); return img

# ---------- PPTX renderer (editable in Canva) ----------
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
def rgb(c): return RGBColor(*c)
def render_pptx(E,W,H,out):
    prs=Presentation(); prs.slide_width=Cm(W); prs.slide_height=Cm(H)
    s=prs.slides.add_slide(prs.slide_layouts[6])
    def shape(kind,x,y,w,h,c,r=None):
        sh=s.shapes.add_shape(kind,Cm(x),Cm(y),Cm(w),Cm(h)); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(c); sh.line.fill.background()
        if r is not None and kind==MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0]=min(0.5,r/min(w,h))
        return sh
    def text(txt,x,y,w,h,fname,size,fill,align,rot,lh,runs=None):
        lsp=(lh/size) if lh else 1.15
        tb=s.shapes.add_textbox(Cm(x),Cm(y-size*0.22),Cm(w),Cm(h+size*0.4)); tf=tb.text_frame; tf.word_wrap=True
        tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=MSO_ANCHOR.TOP
        lines=txt.split('\n')
        for i,l in enumerate(lines):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.alignment=PP_ALIGN.LEFT if align=='left' else PP_ALIGN.CENTER; p.line_spacing=lsp
            if runs and i==0:
                for rt,rc in runs:
                    r=p.add_run(); r.text=rt; r.font.name=PPTX_FONT[fname]; r.font.size=Pt(size/2.54*72); r.font.color.rgb=rgb(rc); r.font.bold=(fname in('nunitob','nunitox'))
            else:
                r=p.add_run(); r.text=l; r.font.name=PPTX_FONT[fname]; r.font.size=Pt(size/2.54*72); r.font.color.rgb=rgb(fill); r.font.bold=(fname in('nunitob','nunitox'))
        if rot: tb.rotation=-rot
    for e in E:
        k=e[0]
        if k=='rect': _,x,y,w,h,c=e; shape(MSO_SHAPE.RECTANGLE,x,y,w,h,c)
        elif k=='roundrect': _,x,y,w,h,c,r=e; shape(MSO_SHAPE.ROUNDED_RECTANGLE,x,y,w,h,c,r)
        elif k=='poly':
            pts=e[1]; xs=[p[0] for p in pts]; ys=[p[1] for p in pts]
            ff=s.shapes.build_freeform(Cm(xs[0]),Cm(ys[0])); ff.add_line_segments([(Cm(a),Cm(b)) for a,b in pts[1:]]+[(Cm(xs[0]),Cm(ys[0]))]); sh=ff.convert_to_shape(); sh.fill.solid(); sh.fill.fore_color.rgb=rgb(e[2]); sh.line.fill.background()
        elif k=='line':
            _,x0,y0,x1,y1,c,wd=e; ln=s.shapes.add_connector(1,Cm(x0),Cm(y0),Cm(x1),Cm(y1)); ln.line.color.rgb=rgb(c); ln.line.width=Pt(wd/2.54*72)
        elif k=='image': _,p,x,y,w,h=e; s.shapes.add_picture(p,Cm(x),Cm(y),Cm(w),Cm(h))
        elif k=='text': text(*e[1:])
        elif k=='text2': _,a,b,x,y,w,h,fname,size,ca,cb=e; text(a+b,x,y,w,h,fname,size,ca,'center',0,None,runs=[(a,ca),(b,cb)])
        elif k=='window':
            _,x,y,w,h=e; sh=shape(MSO_SHAPE.RECTANGLE,x,y,w,h,(255,255,255)); sh.line.fill.solid(); sh.line.color.rgb=rgb((120,120,120)); sh.line.width=Pt(0.75)
    prs.save(out)

# ---------- build ----------
make_assets()
DESIGNS={}
B1=(11.0,15.3); B2=(14.2,21.4); WIN=(0.6,21.4-5.5-4.0,13.0,4.0)
DESIGNS['B1_frente']=(spec_front(*B1),B1); DESIGNS['B1_reverso']=(spec_back(*B1),B1)
DESIGNS['B2_frente']=(spec_front(*B2,window=WIN),B2); DESIGNS['B2_reverso']=(spec_back(*B2,natural=True),B2)
IMGS={}
for name,(E,(W,H)) in DESIGNS.items():
    IMGS[name]=render_pil(E,W,H,f'v3/{name}.png'); render_pptx(E,W,H,f'v3/{name}_editable_canva.pptx'); print(name,'ok')

def sheet(items,out,label,rotate):
    A4=(px(21.0),px(29.7)); sh=Image.new('RGB',A4,'white'); d=ImageDraw.Draw(sh); f=ImageFont.truetype(FONTS['nunito'],30)
    ims=[i.rotate(90,expand=True) if rotate else i for i in items]; gap=px(1.2); total=sum(i.height for i in ims)+gap*(len(ims)-1); y=(A4[1]-total)//2
    for i in ims:
        x=(A4[0]-i.width)//2; sh.paste(i,(x,y))
        for (qx,qy) in [(x,y),(x+i.width,y),(x,y+i.height),(x+i.width,y+i.height)]:
            d.line((qx-70,qy,qx-15,qy),fill='black',width=2); d.line((qx+15,qy,qx+70,qy),fill='black',width=2); d.line((qx,qy-70,qx,qy-15),fill='black',width=2); d.line((qx,qy+15,qx,qy+70),fill='black',width=2)
        y+=i.height+gap
    d.text((px(1),px(0.5)),label+'  ·  imprimir al 100 % (tamaño real, sin ajustar a página)',font=f,fill=(120,120,120)); sh.save(out,resolution=DPI)
sheet([IMGS['B1_frente']]*2,'v3/B1_hoja_frente.pdf','Bolsa 1 (11 x 15.3 cm) · FRENTE · 2 por hoja',True)
sheet([IMGS['B1_reverso']]*2,'v3/B1_hoja_reverso.pdf','Bolsa 1 (11 x 15.3 cm) · REVERSO · 2 por hoja',True)
sheet([IMGS['B2_frente']],'v3/B2_hoja_frente.pdf','Bolsa 2 con ventana (14.2 x 21.4 cm) · FRENTE · 1 por hoja',False)
sheet([IMGS['B2_reverso']],'v3/B2_hoja_reverso.pdf','Bolsa 2 con ventana (14.2 x 21.4 cm) · REVERSO · 1 por hoja',False)
print('sheets ok')
